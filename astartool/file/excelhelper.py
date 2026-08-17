#!/usr/bin/env python
# -*- coding: utf-8 -*-
# @Author: 河北雪域网络科技有限公司 A.Star
# @contact: astar@snowland.ltd
# @site: www.snowland.ltd
# @file: excelhelper.py
# @time: 2019/6/13 17:33
# @Software: PyCharm


__author__ = 'A.Star'

from typing import Union, List, Iterable, Tuple

XLSX_EXTENSIONS = ('.xlsx', '.xlsm', '.xltx', '.xltm')


def _is_xlsx(filename: str) -> bool:
    """
    Judge whether the target file is an xlsx-family workbook by its suffix.

    :param filename: target file path
    :return: ``True`` if the suffix indicates an xlsx-family format
    :rtype: bool
    """
    return filename.lower().endswith(XLSX_EXTENSIONS)


def _write_sheet(rows, booksheet, fields, output_fields, is_dict, writer):
    """
    Write a single sheet into the given workbook sheet object.

    :param rows: data rows (list of dict or list of list/tuple)
    :param booksheet: workbook sheet handle (xlwt Sheet or openpyxl Worksheet)
    :param fields: field names used when ``rows`` is a list of dict
    :param output_fields: header labels written at the first row
    :param is_dict: whether ``rows`` items are dicts
    :param writer: callable ``(sheet, row, col, value)`` to set a cell
    """
    start = 0
    if output_fields:
        for i, field in enumerate(output_fields):
            writer(booksheet, 0, i, field)
        start += 1
    if is_dict:
        for j, obj in enumerate(rows, start):
            for i, field in enumerate(fields):
                writer(booksheet, j, i, obj[field])
    else:
        for j, rowdata in enumerate(rows, start):
            for i, cell in enumerate(rowdata):
                writer(booksheet, j, i, cell)


def to_excel(data,
             filename: str,
             sheetname: Union[str, Union[List[str], Tuple]] = "Sheet1",
             fields: Union[Iterable[str], Iterable[List[str]]] = None, *,
             output_fields: Union[Iterable[str], Iterable[List[str]]] = None, encoding='utf-8'):
    """
    Export data into an Excel file.

    The output format (``.xls`` or ``.xlsx``) is decided by the suffix of
    ``filename``. ``.xls`` is written via the optional ``xlwt`` dependency,
    while ``.xlsx`` is written via the optional ``openpyxl`` dependency. Both
    dependencies are imported lazily and raise a clear ``ImportError`` if
    missing.

    :param data: data to export; a single sheet (list of rows) or multiple
        sheets (list of list of rows) when ``sheetname`` is a collection
    :param filename: output file path; suffix decides the format
    :param sheetname: sheet name, or a collection of names for multi-sheet
    :param fields: field names when ``data`` items are dicts
    :param output_fields: header labels; defaults to ``fields``
    :param encoding: text encoding for the legacy ``.xls`` writer
    :return:
    """
    single = isinstance(sheetname, str)

    if single:
        sheet_specs = [(sheetname, data, fields, output_fields)]
    else:
        if fields is None:
            fields = [None] * len(sheetname)
        if output_fields is None:
            output_fields = fields
        sheet_specs = list(zip(sheetname, data, fields, output_fields))

    if _is_xlsx(filename):
        try:
            import openpyxl
        except ImportError:
            raise ImportError(
                "openpyxl is required to write .xlsx files. Install it via "
                "`pip install openpyxl` or `pip install astartool[optional]`."
            )
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for idx, (it_name, it_data, it_fields, it_output) in enumerate(sheet_specs):
            if idx > 0:
                sheet = workbook.create_sheet(title=it_name)
            else:
                sheet.title = it_name
            is_dict = isinstance(it_data[0], dict)
            if it_output is None:
                it_output = it_fields
            _write_sheet(
                it_data, sheet, it_fields, it_output, is_dict,
                writer=lambda s, r, c, v: s.cell(row=r + 1, column=c + 1, value=v),
            )
        workbook.save(filename)
    else:
        try:
            import xlwt
        except ImportError:
            raise ImportError(
                "xlwt is required to write .xls files. Install it via "
                "`pip install xlwt` or `pip install astartool[optional]`."
            )
        workbook = xlwt.Workbook(encoding=encoding)
        for it_name, it_data, it_fields, it_output in sheet_specs:
            booksheet = workbook.add_sheet(it_name, cell_overwrite_ok=True)
            is_dict = isinstance(it_data[0], dict)
            if it_output is None:
                it_output = it_fields
            _write_sheet(
                it_data, booksheet, it_fields, it_output, is_dict,
                writer=lambda s, r, c, v: s.write(r, c, v),
            )
        workbook.save(filename)
